import shutil
from pathlib import Path

from .data import ZERO_BASED


def group_ids(key, targets):
    return [key[1], *sorted(targets - {key[1]})]


def fasta_text(key, targets, regions, proteins):
    family, source = key
    lines = []
    for lid in group_ids(key, targets):
        region = regions[lid]
        protein = proteins[region.protein_id]
        lines.append(f">fam={family}|src={source}|lcr={lid}|pid={protein.protein_id}"
                     f"|species={protein.species}|start={region.start}|end={region.end}"
                     f"|seg={region.seg_mode}")
        lines.extend(region.sequence[i:i + 80] for i in range(0, len(region.sequence), 80))
    return "\n".join(lines) + "\n"


def export_fastas(directory, groups, regions, proteins):
    directory = Path(directory)
    default, strict = directory / "anchor_fastas_default", directory / "anchor_fastas_strict"
    default.mkdir()
    strict.mkdir()
    strict_count = 0
    for number, (key, targets) in enumerate(sorted(groups.items()), 1):
        if not targets:
            raise ValueError(f"Group {key} has no orthologs")
        name = f"{key[0]}_{key[1]}.fasta"
        path = default / name
        # The supplied Windows FASTAs use CRLF. Keep exports byte-reproducible on every OS.
        path.write_text(fasta_text(key, targets, regions, proteins), encoding="utf-8", newline="\r\n")
        if regions[key[1]].seg_mode == "strict":
            shutil.copyfile(path, strict / name)
            strict_count += 1
        if number % 10000 == 0:
            print(f"FASTA: {number:,}/{len(groups):,} groups", flush=True)
    return {"default_groups": len(groups), "strict_groups": strict_count}


def metadata_tables(groups, regions, proteins, crosswalk=None):
    """Identifiers stay text. Counts/lengths are numeric; missing identifiers stay blank."""
    by_oma, by_acc = crosswalk or ({}, {})
    group_rows, used = [], set()
    for key, targets in sorted(groups.items()):
        fid, src = key
        ids = group_ids(key, targets)
        pids = {regions[lid].protein_id for lid in ids}
        used.update(pids)
        source = regions[src]
        group_rows.append([str(fid), str(src), proteins[source.protein_id].uniprot_acc,
                           source.seg_mode, len(ids),
                           len({proteins[pid].species for pid in pids}), f"{fid}_{src}.fasta"])
    protein_rows = []
    for pid in sorted(used):
        protein = proteins[pid]
        accs = ({protein.uniprot_acc} if protein.uniprot_acc else set()) | by_oma.get(protein.oma_id, set())
        omas = ({protein.oma_id} if protein.oma_id else set()) | by_acc.get(protein.uniprot_acc, set())
        # One-to-many identifier mappings are preserved, never arbitrarily reduced to one.
        protein_rows.append([str(pid), ";".join(sorted(accs)), ";".join(sorted(omas)),
                             protein.species, protein.taxid or None, len(protein.sequence)])
    systems = {regions[lid].coordinate_system for (fid, src), lids in groups.items()
               for lid in (src, *lids)}
    notes = [
        "One FASTA per anchored group. Strict is a subset selected by human source mode.",
        "FASTA files contain extracted sequences, not multiple-sequence alignments.",
    ]
    if ZERO_BASED in systems:
        notes.append("Historical human FASTA coordinates: 0-based half-open; orthologs: 1-based inclusive.")
    else:
        notes.append("FASTA coordinates: 1-based inclusive for all exported regions.")
    return {
        "Groups": {
            "headers": ["family_id", "source_lcr_id", "source_uniprot_acc", "source_mode",
                        "n_sequences", "n_species", "fasta_name"],
            "rows": group_rows,
            "notes": notes,
            "widths": [15, 18, 25, 18, 18, 16, 30],
        },
        "Proteins": {
            "headers": ["protein_id", "uniprot_acc", "oma_id", "species", "taxid", "protein_length_aa"],
            "rows": protein_rows,
            "notes": [
                "Only proteins present in the FASTA outputs. Join using pid in the FASTA header.",
                "Blank IDs/taxids are unavailable in the local inputs. Multiple ID matches use semicolons.",
                "Sources: protein.csv and optional OMA/UniProt crosswalk. Family IDs are local, not OMA group IDs.",
            ],
            "widths": [18, 26, 26, 24, 16, 25],
        },
    }


def write_metadata(path, tables):
    """Portable Windows/local exporter; stream rows to keep memory bounded."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as error:
        raise RuntimeError("XLSX export requires openpyxl: pip install openpyxl") from error
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    workbook.properties.title = "Anchor-based ortholog LCR metadata"
    workbook.properties.creator = "LCR ortholog pipeline"
    for name, table in tables.items():
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "D6" if name == "Proteins" else "C6"
        for number, width in enumerate(table["widths"], 1):
            sheet.column_dimensions[get_column_letter(number)].width = width
        for text in table["notes"]:
            cell = WriteOnlyCell(sheet, text)
            cell.font = Font(name="Arial", size=10, italic=True, color="555555")
            sheet.append([cell])
        sheet.append([])
        headers = []
        for label in table["headers"]:
            cell = WriteOnlyCell(sheet, label)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="263B50")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            headers.append(cell)
        sheet.append(headers)
        for values in table["rows"]:
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10)
                cells.append(cell)
            sheet.append(cells)
        sheet.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{5 + len(table['rows'])}"
    workbook.save(path)
